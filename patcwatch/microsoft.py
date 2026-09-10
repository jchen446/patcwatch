"""Microsoft Graph read-only integration. No browser automation or model client."""
import hashlib
import io
import json
import re
import threading
import time
import uuid
import zipfile
from datetime import datetime, timezone
from html.parser import HTMLParser
from urllib.parse import quote, urlencode, urlsplit

from .scheduler import SourceError
from .server import save_json

SCOPES = ['User.Read', 'Chat.Read', 'Sites.Read.All', 'Files.Read.All']
GRAPH = 'https://graph.microsoft.com/v1.0'


class TrustedHttp:
    """All external HTTP goes through explicit destinations; redirects never inherit auth."""
    def __init__(self):
        import requests
        self.session = requests.Session()
        self.session.trust_env = False
        self.requests = 0
        self.blocked = 0

    def request(self, method, url, *, download_host=None, **kwargs):
        p = urlsplit(url)
        allowed = {'login.microsoftonline.com', 'graph.microsoft.com'}
        if download_host:
            allowed = {download_host}
        if p.scheme != 'https' or p.hostname not in allowed or p.username or p.password or p.port not in (None, 443):
            self.blocked += 1
            raise SourceError('NETWORK_DESTINATION_DENIED')
        if method not in ('GET', 'POST') or (p.hostname == 'graph.microsoft.com' and method != 'GET'):
            raise SourceError('READ_ONLY_POLICY')
        kwargs.pop('allow_redirects', None)
        kwargs.pop('timeout', None)
        kwargs.pop('stream', None)
        self.requests += 1
        import requests
        try:
            response = self.session.request(method, url, timeout=(10, 30), allow_redirects=False, stream=True, **kwargs)
            with response:
                limit = 20_000_000 if download_host else 8_000_000
                chunks, size = [], 0
                for chunk in response.iter_content(65536):
                    size += len(chunk)
                    if size > limit:
                        raise SourceError('RESPONSE_TOO_LARGE')
                    chunks.append(chunk)
                response._content = b''.join(chunks)
                response._content_consumed = True
            return response
        except requests.RequestException:
            raise SourceError('NETWORK_UNAVAILABLE') from None

    def get(self, url, **kwargs):
        return self.request('GET', url, **kwargs)

    def post(self, url, **kwargs):
        return self.request('POST', url, **kwargs)


class Vault:
    def __init__(self, data):
        import keyring
        from cryptography.fernet import Fernet
        backend = keyring.get_keyring()
        # Refuse plaintext, fail, null and arbitrary third-party backends.
        if type(backend).__module__ not in {'keyring.backends.macOS', 'keyring.backends.Windows', 'keyring.backends.SecretService'}:
            raise SourceError('SECURE_KEYCHAIN_UNAVAILABLE')
        account = hashlib.sha256(str(data.resolve()).encode()).hexdigest()
        try:
            key = backend.get_password('Patcwatch', account)
            if not key:
                key = Fernet.generate_key().decode()
                backend.set_password('Patcwatch', account, key)
            self.cipher = Fernet(key.encode())
        except Exception:
            raise SourceError('SECURE_KEYCHAIN_UNAVAILABLE') from None
        self.path = data / 'microsoft-cache.enc'

    def load(self):
        if not self.path.exists():
            return None
        try:
            return self.cipher.decrypt(bytes.fromhex(json.loads(self.path.read_text())['encrypted'])).decode()
        except Exception:
            raise SourceError('TOKEN_CACHE_UNREADABLE') from None

    def save(self, text):
        save_json(self.path, {'encrypted': self.cipher.encrypt(text.encode()).hex()})


class MicrosoftAuth:
    def __init__(self, data, transport):
        self.data, self.http = data, transport
        self.path = data / 'microsoft.json'
        self.config = json.loads(self.path.read_text()) if self.path.exists() else None
        self.lock = threading.RLock()
        self.app = self.cache = self.vault = None
        self.state = 'NOT_CONNECTED'
        self.flow = None
        self.worker = None

    def configure(self, client_id, tenant_id):
        config = {'client_id': str(uuid.UUID(client_id)), 'tenant_id': str(uuid.UUID(tenant_id))}
        with self.lock:
            if self.config and self.config != config:
                raise ValueError('Use a new --data-dir when changing tenant or application.')
            save_json(self.path, config)
            self.config = config

    def initialize(self):
        if self.app:
            return
        if not self.config:
            raise SourceError('APP_REGISTRATION_REQUIRED')
        import msal
        self.vault = Vault(self.data)
        self.cache = msal.SerializableTokenCache()
        saved = self.vault.load()
        if saved:
            self.cache.deserialize(saved)
        self.app = msal.PublicClientApplication(self.config['client_id'],
            authority='https://login.microsoftonline.com/' + self.config['tenant_id'],
            token_cache=self.cache, http_client=self.http, instance_discovery=False)

    def persist(self):
        if self.cache.has_state_changed:
            self.vault.save(self.cache.serialize())

    def start(self):
        with self.lock:
            if self.worker and self.worker.is_alive():
                raise ValueError('A sign-in is already pending.')
            self.initialize()
            self.flow = self.app.initiate_device_flow(scopes=SCOPES)
            if 'user_code' not in self.flow:
                self.state = 'SIGN_IN_DENIED'
                raise SourceError(self.state)
            self.state = 'AWAITING_SIGN_IN'
            public = {'user_code': self.flow['user_code'], 'verification_uri': 'https://microsoft.com/devicelogin', 'expires_in': self.flow.get('expires_in', 900)}
            def acquire():
                try:
                    result = self.app.acquire_token_by_device_flow(self.flow)
                    with self.lock:
                        self.persist()
                        self.state = 'CONNECTED' if result.get('access_token') else 'SIGN_IN_REQUIRED'
                except Exception:
                    with self.lock:
                        self.state = 'SIGN_IN_FAILED'
                finally:
                    with self.lock:
                        self.flow = None
            self.worker = threading.Thread(target=acquire, daemon=True)
            self.worker.start()
            return public

    def token(self, force=False):
        with self.lock:
            if self.state == 'AWAITING_SIGN_IN':
                raise SourceError('SIGN_IN_REQUIRED')
            self.initialize()
            accounts = self.app.get_accounts()
            if len(accounts) != 1:
                self.state = 'SIGN_IN_REQUIRED'
                raise SourceError(self.state)
            result = self.app.acquire_token_silent_with_error(SCOPES, account=accounts[0], force_refresh=force)
            self.persist()
            if not result or not result.get('access_token'):
                self.state = 'SIGN_IN_REQUIRED'
                raise SourceError(self.state)
            self.state = 'CONNECTED'
            return result['access_token']

    def status(self):
        return {'state': self.state, 'configured': bool(self.config), 'config': self.config, 'scopes': SCOPES}


class Text(HTMLParser):
    def __init__(self):
        super().__init__(); self.parts = []
    def handle_data(self, data):
        self.parts.append(data)
    def handle_starttag(self, tag, attrs):
        if tag in ('br', 'p', 'div', 'li'): self.parts.append('\n')


def text_content(raw):
    p = Text(); p.feed(raw); return ''.join(p.parts).strip()


class Graph:
    def __init__(self, auth, http):
        self.auth, self.http = auth, http

    def get(self, path):
        url = path if path.startswith('https://') else GRAPH + path
        p = urlsplit(url)
        if p.scheme != 'https' or p.netloc != 'graph.microsoft.com' or not p.path.startswith('/v1.0/'):
            raise SourceError('PAGINATION_DESTINATION_DENIED')
        for attempt in range(2):
            token = self.auth.token(force=attempt == 1)
            response = self.http.get(url, headers={'Authorization': 'Bearer ' + token})
            if response.status_code != 401:
                break
        if response.status_code == 401:
            raise SourceError('SIGN_IN_REQUIRED')
        if response.status_code == 403:
            raise SourceError('PERMISSION_DENIED')
        if response.status_code in (429, 503):
            retry = response.headers.get('Retry-After', '60')
            try: retry = min(86400, max(1, float(retry)))
            except ValueError: retry = 60
            raise SourceError('THROTTLED' if response.status_code == 429 else 'SERVICE_UNAVAILABLE', retry)
        if response.status_code != 200:
            raise SourceError('SOURCE_UNAVAILABLE')
        try:
            value = response.json()
            if not isinstance(value, dict): raise ValueError()
            return value
        except ValueError:
            raise SourceError('INVALID_RESPONSE') from None

    def pages(self, path, max_pages=200):
        result, visited = [], set()
        while path:
            if path in visited or len(visited) >= max_pages:
                raise SourceError('INCOMPLETE_PAGINATION')
            visited.add(path)
            value = self.get(path)
            rows = value.get('value')
            if not isinstance(rows, list) or any(not isinstance(r, dict) for r in rows):
                raise SourceError('INVALID_RESPONSE')
            result.extend(rows)
            if len(result) > 20000:
                raise SourceError('INCOMPLETE_PAGINATION')
            path = value.get('@odata.nextLink')
            if path is not None and not isinstance(path, str):
                raise SourceError('INVALID_RESPONSE')
        return result

    def chats(self, query):
        all_chats = self.pages('/me/chats?$top=50')
        needle = query.casefold().strip()
        if not needle: raise ValueError('Enter a chat name')
        return [{'id': c['id'], 'name': c.get('topic') or '(Unnamed chat)', 'type': c.get('chatType', '')}
                for c in all_chats if needle in (c.get('topic') or '').casefold()]

    def site(self, url):
        p = urlsplit(url)
        if p.scheme != 'https' or not p.hostname or not p.hostname.endswith('.sharepoint.com') or p.port not in (None,443) or p.username or p.query or p.fragment:
            raise ValueError('Enter a SharePoint site URL, for example https://tenant.sharepoint.com/sites/QA')
        return self.get('/sites/' + p.hostname + ':' + quote(p.path.rstrip('/') or '/', safe='/'))

    def drives(self, site_id):
        return self.pages('/sites/' + quote(site_id, safe='') + '/drives')

    def children(self, drive_id, folder_id='root'):
        base = '/drives/' + quote(drive_id, safe='')
        return self.pages(base + ('/root/children' if folder_id == 'root' else '/items/' + quote(folder_id,safe='') + '/children'))

    def workbook(self, drive_id, item_id, host):
        base = '/drives/' + quote(drive_id, safe='') + '/items/' + quote(item_id, safe='')
        meta = self.get(base)
        if not meta.get('name', '').lower().endswith('.xlsx') or meta.get('size', 0) > 20_000_000:
            raise SourceError('UNSUPPORTED_WORKBOOK')
        if urlsplit(meta.get('webUrl', '')).hostname != host:
            raise SourceError('SOURCE_IDENTITY_MISMATCH')
        url = meta.get('@microsoft.graph.downloadUrl')
        if not isinstance(url, str): raise SourceError('DOWNLOAD_UNAVAILABLE')
        # A preauthenticated URL is a secret. Never persist it or send an auth header.
        response = self.http.request('GET', url, download_host=host)
        if response.status_code != 200:
            raise SourceError('DOWNLOAD_UNAVAILABLE')
        after = self.get(base)
        if not meta.get('eTag') or meta['eTag'] != after.get('eTag'):
            raise SourceError('SOURCE_CHANGED_DURING_READ')
        return decode_workbook(response.content)


def decode_workbook(raw):
    import openpyxl
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if len(z.infolist()) > 1000 or sum(i.file_size for i in z.infolist()) > 100_000_000:
                raise SourceError('WORKBOOK_TOO_LARGE')
        book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        sheets = {}
        try:
            if len(book.sheetnames) > 30: raise SourceError('WORKBOOK_TOO_LARGE')
            for sheet in book:
                if sheet.max_row and sheet.max_row > 20001 or sheet.max_column and sheet.max_column > 100:
                    raise SourceError('WORKBOOK_TOO_LARGE')
                rows = []
                for row in sheet.iter_rows():
                    if len(rows) >= 20001 or len(row) > 100: raise SourceError('WORKBOOK_TOO_LARGE')
                    values = ['' if c.value is None else str(c.value) for c in row]
                    if any(len(v) > 10000 for v in values): raise SourceError('WORKBOOK_TOO_LARGE')
                    rows.append(values)
                sheets[sheet.title] = rows
        finally:
            book.close()
        return sheets
    except SourceError: raise
    except Exception: raise SourceError('INVALID_WORKBOOK') from None


def workbook_items(sheets, config):
    rows = sheets.get(config['sheet'])
    header_row = config['header_row']
    if not rows or not isinstance(header_row, int) or header_row < 1 or header_row > len(rows):
        raise SourceError('ROW_MAPPING_UNAVAILABLE')
    headers = rows[header_row - 1]
    needed = [config['id_column'], *config['columns']]
    if any(headers.count(h) != 1 for h in needed) or not config['columns']:
        raise SourceError('ROW_MAPPING_UNAVAILABLE')
    id_index = headers.index(config['id_column'])
    indexes = [(h, headers.index(h)) for h in config['columns']]
    items, seen = [], set()
    for row in rows[header_row:]:
        if not any(row): continue
        if len(row) <= max([id_index, *[i for h,i in indexes]]): raise SourceError('ROW_MAPPING_UNAVAILABLE')
        ident = row[id_index].strip()
        if not ident or ident in seen: raise SourceError('DUPLICATE_OR_MISSING_ROW_ID')
        seen.add(ident)
        items.append({'id': ident, 'content': json.dumps({h:row[i] for h,i in indexes}, ensure_ascii=False, sort_keys=True)})
    return items


class Collector:
    def __init__(self, graph, clock=time.time):
        self.graph, self.clock = graph, clock

    def read(self, source):
        c = source['config']
        if source['kind'] == 'teams':
            start = source['cursor'].get('last_poll', self.clock() - 7*86400)
            # Overlap accounts for messages edited while polling; native IDs deduplicate.
            stamp = datetime.fromtimestamp(start - 300, timezone.utc).isoformat().replace('+00:00','Z')
            begun = self.clock()
            query = urlencode({'$top':50, '$orderby':'lastModifiedDateTime desc', '$filter':'lastModifiedDateTime gt ' + stamp})
            rows = self.graph.pages('/chats/' + quote(c['chat_id'],safe='') + '/messages?' + query)
            items = []
            for r in rows:
                if not isinstance(r.get('id'), str) or not isinstance(r.get('body'), dict):
                    raise SourceError('INVALID_RESPONSE')
                body = r['body'].get('content', '')
                content = {'text':text_content(body) if r['body'].get('contentType', 'html') == 'html' else body, 'deleted':bool(r.get('deletedDateTime')),
                           'attachments':[{'id':a.get('id'), 'name':a.get('name'), 'contentType':a.get('contentType')} for a in r.get('attachments',[])]}
                items.append({'id':r['id'],'content':json.dumps(content,ensure_ascii=False,sort_keys=True)})
            return {'complete':True,'items':items,'cursor':{'last_poll':begun},
                    'coverage':'Messages modified since previous poll with five-minute overlap; initial baseline seven days. Attachment bodies not read.'}
        if source['kind'] == 'sharepoint':
            sheets = self.graph.workbook(c['drive_id'], c['item_id'], c['host'])
            items = workbook_items(sheets,c)
            return {'complete':True,'items':items,'coverage':'Selected worksheet and mapped columns, stable row IDs; cached formula values only.'}
        raise SourceError('UNSUPPORTED_SOURCE')
